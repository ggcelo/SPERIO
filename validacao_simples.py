from io import BytesIO

import streamlit as st
from openpyxl import load_workbook


NOME_ABA = "Gabarito e validação"


def ler_planilha(conteudo_planilha):
    """Lê a planilha e organiza os casos pelo código."""

    planilha = load_workbook(BytesIO(conteudo_planilha), data_only=False)
    aba = planilha[NOME_ABA]
    cabecalhos = {}

    for numero_coluna, celula in enumerate(aba[1], start=1):
        if celula.value is not None:
            cabecalhos[celula.value] = numero_coluna

    casos = {}

    for numero_linha in range(2, aba.max_row + 1):
        codigo = aba.cell(numero_linha, cabecalhos["id_caso"]).value

        if codigo is None:
            continue

        codigo = str(codigo).strip().upper()
        casos[codigo] = {}

        for nome_coluna, numero_coluna in cabecalhos.items():
            casos[codigo][nome_coluna] = aba.cell(numero_linha, numero_coluna).value

    return casos


def importar_caso(caso):
    """Preenche os dicionários utilizados pelos módulos do software."""

    # Trabalha com uma cópia para não alterar os dados lidos da planilha.
    caso = caso.copy()

    textos_ausentes = [
        "",
        "AUS",
        "Ausente",
        "NA",
        "Não se aplica",
        "NC",
        "Não coletado",
        "NI",
        "Não informado"
    ]

    # Troca os códigos de ausência por None.
    for campo, valor in caso.items():
        if valor is None:
            caso[campo] = None

        elif isinstance(valor, str) and valor.strip() in textos_ausentes:
            caso[campo] = None

    # Converte somente os campos que precisam ser numéricos.
    campos_decimais = [
        "peso_kg",
        "altura_m",
        "cintura_cm",
        "quadril_cm",
        "pbl_max_pct",
        "cal_max_mm",
        "sondagem_max_mm",
        "perda_ossea_vertical_max_mm",
        "percentual_dentes_envolvidos",
        "progressao_5_anos_mm",
        "hba1c_pct"
    ]

    campos_inteiros = [
        "idade_anos",
        "cigarros_dia",
        "dentes_perdidos_periodontite",
        "dentes_remanescentes"
    ]

    for campo in campos_decimais:
        if caso[campo] is not None:
            caso[campo] = float(caso[campo])

    for campo in campos_inteiros:
        if caso[campo] is not None:
            caso[campo] = int(float(caso[campo]))

    imc = None

    if caso["peso_kg"] is not None and caso["altura_m"] is not None:
        imc = round(caso["peso_kg"] / (caso["altura_m"] ** 2), 2)

    # Informações compartilhadas.
    st.session_state["dados_caso"] = {
        "codigo": caso["id_caso"],
        "idade": caso["idade_anos"],
        "sexo": caso["sexo"]
    }

    st.session_state["dados_clinicos"] = {
        "peso": caso["peso_kg"],
        "altura": caso["altura_m"],
        "cintura": caso["cintura_cm"],
        "quadril": caso["quadril_cm"],
        "imc": imc,
        "diabetes": caso["diabetes_diagnosticado"],
        "tabagismo": caso["tabagismo"],
        "cigarros_dia": caso["cigarros_dia"]
    }

    # Dados do FINDRISC.
    st.session_state["dados_findrisc"] = {
        "atividade_fisica": caso["atividade_30min_dia"],
        "frutas_vegetais": caso["frutas_vegetais_diarios"],
        "medicacao_pressao": caso["uso_medicamento_pressao"],
        "glicemia_elevada": caso["glicemia_elevada_previa"],
        "historico_familiar": caso["historico_familiar_diabetes"]
    }

    # Dados do INTERHEART.
    st.session_state["dados_interheart"] = {
        "hipertensao": caso["hipertensao_diagnosticada"],
        "historico_parental_infarto": caso[
            "historico_parental_infarto"
        ],
        "estresse": caso["estresse"],
        "fumo_passivo": caso["fumo_passivo_1h_semana"],
        "atividade_lazer": caso["atividade_fisica_lazer"],
        "tristeza_depressao": caso[
            "tristeza_depressao_2_semanas"
        ],
        "alimento_salgado": caso["alimento_salgado_diario"],
        "fritura_fast_food": caso[
            "fritura_fast_food_3x_semana"
        ],
        "fruta_diaria": caso["fruta_diaria"],
        "vegetais_diarios": caso["vegetais_diarios"],
        "carne_aves": caso["carne_aves_2x_dia"]
    }

    # Dados periodontais.
    st.session_state["dados_periodontais"] = {
        "cal": caso["cal_max_mm"],
        "dentes_perdidos": caso["dentes_perdidos_periodontite"],
        "sondagem": caso["sondagem_max_mm"],
        "perda_vertical": caso["perda_ossea_vertical_max_mm"],
        "furca": caso["furca"],
        "defeito_rebordo": caso["defeito_rebordo"],
        "disfuncao_mastigatoria": caso["disfuncao_mastigatoria"],
        "mobilidade": caso["mobilidade"],
        "dentes_remanescentes": caso["dentes_remanescentes"],
        "colapso_mordida": caso["colapso_mordida"],
        "migracao_dentaria": caso["migracao_dentaria"],
        "abertura_leque": caso["abertura_leque"],
        "percentual_dentes_envolvidos": caso[
            "percentual_dentes_envolvidos"
        ],
        "padrao_molar_incisivo": caso["padrao_molar_incisivo"],
        "progressao_5_anos": caso["progressao_5_anos_mm"],
        "fenotipo": caso["fenotipo_biofilme"],
        "hba1c": caso["hba1c_pct"]
    }

    # O %PBL é colocado diretamente no estado apenas para a validação.
    pbl = caso["pbl_max_pct"]

    st.session_state["avaliacao"] = {
        "periodontal": None,
        "findrisc": None,
        "interheart": None,
        "pbl": pbl
    }

    st.session_state["dados_radiograficos"] = {
        "nome_arquivo": None,
        "imagem_bytes": None,
        "resultados_dentes": [],
        "pbl_maximo": pbl,
        "dente_determinante": "Caso sintético" if pbl is not None else None,
        "superficie_determinante": (
            "Valor da planilha" if pbl is not None else None
        ),
        "confirmado": pbl is not None,
        "imagem_anotada_bytes": None,
        "pbl_calculado": pbl is not None
    }


def salvar_resultados(conteudo_planilha, codigo, observacao):
    """Grava os resultados do software nas colunas amarelas."""

    periodontal = st.session_state["avaliacao"]["periodontal"]
    findrisc = st.session_state["avaliacao"]["findrisc"]
    interheart = st.session_state["avaliacao"]["interheart"]
    diabetes = st.session_state["dados_clinicos"]["diabetes"]

    if periodontal is None:
        return None, "A classificação periodontal ainda não foi executada."

    if diabetes != "Sim" and (findrisc is None or "pontos" not in findrisc):
        return None, "O FINDRISC ainda não foi executado."

    if interheart is None:
        return None, "O INTERHEART ainda não foi executado."

    estagio = periodontal.get("estagio") or "NC"
    extensao = periodontal.get("extensao_distribuicao")
    grau = periodontal.get("grau") or "NC"

    if extensao is None:
        extensao = "NC"

    elif extensao == "Padrão molar/incisivo":
        extensao = "padrão molar-incisivo"

    else:
        extensao = extensao.lower()

    if diabetes == "Sim":
        pontos_findrisc = "NA"
        classificacao_findrisc = "Não aplicável"

    else:
        pontos_findrisc = findrisc["pontos"]
        classificacao_findrisc = findrisc["classificacao"]

    if "pontos" in interheart:
        pontos_interheart = interheart["pontos"]
        classificacao_interheart = interheart["classificacao"]

    else:
        pontos_interheart = "NC"
        classificacao_interheart = "NC"

    planilha = load_workbook(BytesIO(conteudo_planilha), data_only=False)
    aba = planilha[NOME_ABA]
    cabecalhos = {}

    for numero_coluna, celula in enumerate(aba[1], start=1):
        if celula.value is not None:
            cabecalhos[celula.value] = numero_coluna

    linha_do_caso = None

    for numero_linha in range(2, aba.max_row + 1):
        codigo_da_linha = aba.cell(numero_linha, cabecalhos["id_caso"]).value

        if str(codigo_da_linha).strip().upper() == codigo.upper():
            linha_do_caso = numero_linha
            break

    if linha_do_caso is None:
        return None, "O caso não foi encontrado na planilha."

    resultados = {
        "estagio_software": estagio,
        "extensao_software": extensao,
        "grau_software": grau,
        "findrisc_pontos_software": pontos_findrisc,
        "findrisc_classificacao_software": classificacao_findrisc,
        "interheart_pontos_software": pontos_interheart,
        "interheart_classificacao_software": classificacao_interheart
    }

    for nome_coluna, valor in resultados.items():
        aba.cell(linha_do_caso, cabecalhos[nome_coluna]).value = valor

    aba.cell(linha_do_caso, cabecalhos["observacao"]).value = observacao or None 
    arquivo_atualizado = BytesIO()
    planilha.save(arquivo_atualizado)

    return arquivo_atualizado.getvalue(), None


def montar_dados_paciente_exportacao():
    """Reúne os dados do paciente usando os cabeçalhos do gabarito."""

    dados_caso = st.session_state["dados_caso"]
    dados_clinicos = st.session_state["dados_clinicos"]
    dados_findrisc = st.session_state["dados_findrisc"]
    dados_interheart = st.session_state["dados_interheart"]
    dados_periodontais = st.session_state["dados_periodontais"]
    avaliacao = st.session_state["avaliacao"]

    resultado_periodontal = avaliacao["periodontal"] or {}
    resultado_findrisc = avaliacao["findrisc"] or {}
    resultado_interheart = avaliacao["interheart"] or {}

    extensao = resultado_periodontal.get("extensao_distribuicao")

    if extensao == "Padrão molar/incisivo":
        extensao = "padrão molar-incisivo"
    elif extensao is not None:
        extensao = extensao.lower()

    pontos_findrisc = resultado_findrisc.get("pontos")
    classificacao_findrisc = resultado_findrisc.get("classificacao")

    if dados_clinicos["diabetes"] == "Sim":
        pontos_findrisc = "NA"
        classificacao_findrisc = "Não aplicável"

    dados_paciente = {
        "id_caso": dados_caso["codigo"],
        "idade_anos": dados_caso["idade"],
        "sexo": dados_caso["sexo"],
        "peso_kg": dados_clinicos["peso"],
        "altura_m": dados_clinicos["altura"],
        "cintura_cm": dados_clinicos["cintura"],
        "quadril_cm": dados_clinicos["quadril"],
        "diabetes_diagnosticado": dados_clinicos["diabetes"],
        "tabagismo": dados_clinicos["tabagismo"],
        "cigarros_dia": dados_clinicos["cigarros_dia"],
        "atividade_30min_dia": dados_findrisc["atividade_fisica"],
        "frutas_vegetais_diarios": dados_findrisc["frutas_vegetais"],
        "uso_medicamento_pressao": dados_findrisc["medicacao_pressao"],
        "glicemia_elevada_previa": dados_findrisc["glicemia_elevada"],
        "historico_familiar_diabetes": dados_findrisc[
            "historico_familiar"
        ],
        "hipertensao_diagnosticada": dados_interheart["hipertensao"],
        "historico_parental_infarto": dados_interheart[
            "historico_parental_infarto"
        ],
        "estresse": dados_interheart["estresse"],
        "fumo_passivo_1h_semana": dados_interheart["fumo_passivo"],
        "atividade_fisica_lazer": dados_interheart["atividade_lazer"],
        "tristeza_depressao_2_semanas": dados_interheart[
            "tristeza_depressao"
        ],
        "alimento_salgado_diario": dados_interheart["alimento_salgado"],
        "fritura_fast_food_3x_semana": dados_interheart[
            "fritura_fast_food"
        ],
        "fruta_diaria": dados_interheart["fruta_diaria"],
        "vegetais_diarios": dados_interheart["vegetais_diarios"],
        "carne_aves_2x_dia": dados_interheart["carne_aves"],
        "pbl_max_pct": avaliacao["pbl"],
        "cal_max_mm": dados_periodontais["cal"],
        "dentes_perdidos_periodontite": dados_periodontais[
            "dentes_perdidos"
        ],
        "sondagem_max_mm": dados_periodontais["sondagem"],
        "perda_ossea_vertical_max_mm": dados_periodontais["perda_vertical"],
        "furca": dados_periodontais["furca"],
        "defeito_rebordo": dados_periodontais["defeito_rebordo"],
        "disfuncao_mastigatoria": dados_periodontais[
            "disfuncao_mastigatoria"
        ],
        "mobilidade": dados_periodontais["mobilidade"],
        "dentes_remanescentes": dados_periodontais["dentes_remanescentes"],
        "colapso_mordida": dados_periodontais["colapso_mordida"],
        "migracao_dentaria": dados_periodontais["migracao_dentaria"],
        "abertura_leque": dados_periodontais["abertura_leque"],
        "percentual_dentes_envolvidos": dados_periodontais[
            "percentual_dentes_envolvidos"
        ],
        "padrao_molar_incisivo": dados_periodontais[
            "padrao_molar_incisivo"
        ],
        "progressao_5_anos_mm": dados_periodontais["progressao_5_anos"],
        "fenotipo_biofilme": dados_periodontais["fenotipo"],
        "hba1c_pct": dados_periodontais["hba1c"],
        "estagio_software": resultado_periodontal.get("estagio"),
        "extensao_software": extensao,
        "grau_software": resultado_periodontal.get("grau"),
        "findrisc_pontos_software": pontos_findrisc,
        "findrisc_classificacao_software": classificacao_findrisc,
        "interheart_pontos_software": resultado_interheart.get("pontos"),
        "interheart_classificacao_software": resultado_interheart.get(
            "classificacao"
        ),
        "observacao": None
    }

    return dados_paciente


def adicionar_paciente_planilha(conteudo_planilha):
    """Acrescenta o paciente atual ao final da planilha-gabarito."""

    try:
        planilha = load_workbook(
            BytesIO(conteudo_planilha),
            data_only=False
        )
    except Exception:
        return None, "Não foi possível abrir a planilha."

    if NOME_ABA not in planilha.sheetnames:
        return None, f'A aba "{NOME_ABA}" não foi encontrada.'

    aba = planilha[NOME_ABA]
    cabecalhos = [celula.value for celula in aba[1]]
    dados_paciente = montar_dados_paciente_exportacao()

    cabecalhos_ausentes = [
        nome_coluna
        for nome_coluna in dados_paciente
        if nome_coluna not in cabecalhos
    ]

    if len(cabecalhos_ausentes) > 0:
        return (
            None,
            "A planilha não possui todos os cabeçalhos esperados: "
            + ", ".join(cabecalhos_ausentes)
        )

    codigo = dados_paciente["id_caso"]

    if codigo is None or str(codigo).strip() == "":
        return None, "Informe e salve o código do caso antes de exportar."

    coluna_codigo = cabecalhos.index("id_caso") + 1

    for numero_linha in range(2, aba.max_row + 1):
        codigo_existente = aba.cell(numero_linha, coluna_codigo).value

        if (
            codigo_existente is not None
            and str(codigo_existente).strip().upper()
            == str(codigo).strip().upper()
        ):
            return None, f"O caso {codigo} já existe na planilha."

    nova_linha = [
        dados_paciente.get(nome_coluna)
        for nome_coluna in cabecalhos
    ]

    aba.append(nova_linha)

    arquivo_atualizado = BytesIO()
    planilha.save(arquivo_atualizado)

    return arquivo_atualizado.getvalue(), None


def mostrar_validacao():
    """Mostra o painel de importação e gravação."""

    st.caption(
        "Carregue a planilha, digite o código e importe o caso para todos "
        "os módulos."
    )

    arquivo = st.file_uploader(
        "Planilha-gabarito (.xlsx)",
        type=["xlsx"],
        key="arquivo_validacao"
    )

    if arquivo is not None:
        assinatura_arquivo = arquivo.name + str(arquivo.size)

        if (
            st.session_state.get("assinatura_planilha_validacao")
            != assinatura_arquivo
        ):
            st.session_state["planilha_validacao"] = arquivo.getvalue()
            st.session_state["nome_planilha_validacao"] = arquivo.name
            st.session_state["assinatura_planilha_validacao"] = (
                assinatura_arquivo
            )
            st.session_state.pop("codigo_caso_validacao", None)

    if "planilha_validacao" not in st.session_state:
        return

    try:
        casos = ler_planilha(st.session_state["planilha_validacao"])

    except Exception:
        st.error("Não foi possível ler a planilha-gabarito.")
        return

    codigo_digitado = st.text_input("Código do caso", placeholder="Exemplo: C07").strip().upper()

    if st.button("Importar caso", type="primary", width="stretch"):
        if codigo_digitado == "":
            st.warning("Digite o código do caso.")

        elif codigo_digitado not in casos:
            st.error("O código informado não existe na planilha.")

        else:
            importar_caso(casos[codigo_digitado])
            st.session_state["codigo_caso_validacao"] = codigo_digitado
            st.session_state["modulo_navegacao"] = (
                "Classificação periodontal"
            )
            st.session_state["mensagem_importacao"] = (
                f"Caso {codigo_digitado} importado para todos os módulos."
            )
            st.rerun()

    if "mensagem_importacao" in st.session_state:
        st.success(st.session_state.pop("mensagem_importacao"))

    if "codigo_caso_validacao" not in st.session_state:
        return

    codigo_atual = st.session_state["codigo_caso_validacao"]
    st.write(f"**Caso ativo:** {codigo_atual}")

    observacao = st.text_input("Observação opcional", key="observacao_validacao")

    if st.button("Registrar resultados na planilha", width="stretch"):
        arquivo_atualizado, erro = salvar_resultados(
            st.session_state["planilha_validacao"],
            codigo_atual,
            observacao
        )

        if erro is not None:
            st.warning(erro)

        else:
            st.session_state["planilha_validacao"] = arquivo_atualizado
            st.success(f"Resultados do caso {codigo_atual} registrados.")

    nome_original = st.session_state["nome_planilha_validacao"]
    nome_saida = nome_original.replace(".xlsx", "_resultados.xlsx")

    st.download_button(
        "Baixar planilha com os resultados",
        data=st.session_state["planilha_validacao"],
        file_name=nome_saida,
        mime=("application/vnd.openxmlformats-officedocument.""spreadsheetml.sheet"),
        width="stretch"
    )
