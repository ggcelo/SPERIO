"""Confere os cálculos reais do aplicativo contra os 50 casos do gabarito."""

from io import BytesIO
from pathlib import Path
import unittest

from openpyxl import load_workbook

import modulo_findrisc
import modulo_interheart
import modulo_periodontal
import planilhas


class StreamlitSimulado:
    def __init__(self):
        self.session_state = {}

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False

    def columns(self, quantidade):
        return [self] * quantidade

    def form(self, *_args, **_kwargs):
        return self

    def expander(self, *_args, **_kwargs):
        return self

    def form_submit_button(self, *_args, **_kwargs):
        return True

    def number_input(self, *_args, **kwargs):
        return kwargs["value"]

    def radio(self, _label, opcoes, **kwargs):
        indice = kwargs.get("index")
        return opcoes[indice] if indice is not None else None

    selectbox = radio

    def __getattr__(self, _nome):
        return lambda *_args, **_kwargs: None


class ValidacaoPlanilha(unittest.TestCase):
    def test_todos_os_campos_dos_50_casos(self):
        caminho = Path(__file__).resolve().parents[1] / (
            "planilha_gabarito_50_casos_resultados_logica_atual.xlsx"
        )
        conteudo = caminho.read_bytes()
        casos = planilhas.ler_planilha(conteudo)
        self.assertEqual(len(casos), 50)
        campos = (
            "estagio",
            "extensao",
            "grau",
            "findrisc_pontos",
            "findrisc_classificacao",
            "interheart_pontos",
            "interheart_classificacao",
        )

        # Limpa resultados antigos para exigir que cada campo seja recalculado.
        planilha = load_workbook(BytesIO(conteudo))
        aba = planilha.active
        cabecalhos = {celula.value: i for i, celula in enumerate(aba[1], 1)}
        for linha in range(2, aba.max_row + 1):
            if aba.cell(linha, cabecalhos["id_caso"]).value is None:
                continue
            for campo in campos:
                aba.cell(linha, cabecalhos[f"{campo}_software"]).value = None
        arquivo_limpo = BytesIO()
        planilha.save(arquivo_limpo)
        conteudo = arquivo_limpo.getvalue()

        interface = StreamlitSimulado()
        for modulo in (
            planilhas,
            modulo_periodontal,
            modulo_findrisc,
            modulo_interheart,
        ):
            modulo.st = interface

        for codigo, caso in casos.items():
            planilhas.importar_caso(caso)
            modulo_periodontal.mostrar_modulo_periodontal()
            modulo_findrisc.mostrar_findrisc()
            modulo_interheart.mostrar_interheart()
            conteudo, erro = planilhas.salvar_resultados(conteudo, codigo, "")
            self.assertIsNone(erro, codigo)

        aba = load_workbook(BytesIO(conteudo), data_only=True).active
        cabecalhos = {celula.value: i for i, celula in enumerate(aba[1])}
        for linha in aba.iter_rows(min_row=2, values_only=True):
            codigo = linha[cabecalhos["id_caso"]]
            if codigo is None:
                continue
            for campo in campos:
                with self.subTest(caso=codigo, campo=campo):
                    self.assertEqual(
                        linha[cabecalhos[f"{campo}_software"]],
                        linha[cabecalhos[f"{campo}_gabarito"]],
                    )


if __name__ == "__main__":
    unittest.main()
